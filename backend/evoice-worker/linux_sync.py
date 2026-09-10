#!/usr/bin/env python3
"""Linux eVoice sync: docs/ → audios/*.mp3 via Piper or espeak-ng + ffmpeg.

Adapted for eduardoos.com VPS media layout (local filesystem, not S3):
  media/evoice/<userId>/<project>/{docs,audios}/

Usage (Go API shells this out when EVOICE_FAKE_TTS is unset/false):
  python3 linux_sync.py --project-dir /path/to/project [--only name] [--mode standard|premium|super_premium]

Environment (documented for operators; Go sets most of these):
  EVOICE_PYTHON          — Python interpreter (default: python3); used by Go, not this script
  EVOICE_WORKER_SCRIPT   — Absolute path to this file
  EVOICE_FAKE_TTS        — When true/1, Go uses FakeRunner and never calls this script
  EVOICE_MEDIA_ROOT      — VPS root for evoice files (Go); typically MEDIA_ROOT/evoice
  EVOICE_JOB_TIMEOUT     — Go convert timeout (e.g. 45m)
  EVOICE_WORK_DIR        — Optional temp base (legacy); jobs now run in-place under media
  DEEPSEEK_API_KEY       — Premium / super_premium DeepSeek refine + vision
  TMPDIR / TEMP / TMP    — Set by Go to the parent of --project-dir

Emits frequent progress lines and a final STATS line for the Go job poller.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

DOC_EXTENSIONS = {
    ".docx",
    ".txt",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp", ".gif"}

PREMIUM_SYSTEM = (
    "Eres un editor de guiones hablados en español para audiolibros / MP3. "
    "El texto de usuario puede venir de cualquier modalidad (texto pegado, .txt, .docx, "
    "PDF con capa de texto, PDF escaneado/OCR/visión, o imagen). "
    "Tu trabajo es la preparación perfecta para síntesis de voz (TTS): oraciones cortas y claras, "
    "expande abreviaturas, quita markdown/URLs/ruido OCR, corrige artefactos de extracción, "
    "mantén el significado y el orden lógico. "
    "Divide el contenido en capítulos lógicos (introducción, secciones temáticas, cierre). "
    "Responde SOLO con capítulos en este formato exacto (sin texto fuera de los bloques):\n"
    '<<<CHAPTER n="1" title="Título corto">>>\n'
    "texto hablado del capítulo...\n"
    "<<<END>>>\n"
    '<<<CHAPTER n="2" title="Otro título">>>\n'
    "...\n"
    "<<<END>>>"
)

VISION_PAGE_PROMPT = (
    "Extrae TODO el texto legible de esta página de documento. "
    "Conserva el orden de lectura. No inventes contenido. "
    "Responde solo con el texto plano de la página."
)

# Below this, treat PDF text-layer as sparse and fall back to page OCR.
PDF_TEXT_MIN_CHARS = 40
VISION_DPI = 200
# JPEG keeps Vision payloads smaller than PNG and reduces OOM risk on long PDFs.
VISION_JPEG_QUALITY = 85
VISION_CHECKPOINT_EVERY = 5
VERSION_AUDIO_RE = re.compile(
    r"^(?P<stem>.+)\.v(?P<ver>\d+)(?:\.c\d+-.*)?\.mp3$",
    re.IGNORECASE,
)

CHAPTER_RE = re.compile(
    r'<<<CHAPTER\s+n="(?P<n>\d+)"\s+title="(?P<title>[^"]*)"\s*>>>\s*(?P<body>.*?)\s*<<<END>>>',
    re.IGNORECASE | re.DOTALL,
)


def log(msg: str) -> None:
    print(msg, flush=True)


def needs_regen(doc: Path, mp3: Path) -> bool:
    if not mp3.is_file():
        return True
    return doc.stat().st_mtime > mp3.stat().st_mtime


def slug_title(title: str, max_len: int = 40) -> str:
    raw = (title or "").strip().lower()
    out: list[str] = []
    for ch in raw:
        if ch.isalnum() and ord(ch) < 128:
            out.append(ch)
        elif ch in " -_":
            out.append("-")
    slug = "".join(out)
    while "--" in slug:
        slug = slug.replace("--", "-")
    slug = slug.strip("-") or "cap"
    return slug[:max_len]


def parse_chapters(text: str) -> list[tuple[int, str, str]]:
    """Return list of (n, title, body). Fallback: single chapter if markers missing."""
    found: list[tuple[int, str, str]] = []
    for m in CHAPTER_RE.finditer(text or ""):
        n = int(m.group("n"))
        title = (m.group("title") or f"Capítulo {n}").strip()
        body = (m.group("body") or "").strip()
        if body:
            found.append((n, title, body))
    if found:
        found.sort(key=lambda x: x[0])
        return found
    stripped = (text or "").strip()
    if not stripped:
        return []
    return [(1, "Completo", stripped)]


def chapter_mp3_name(stem: str, n: int, title: str, version: int = 0) -> str:
    slug = slug_title(title)
    if version and version > 0:
        return f"{stem}.v{version}.c{n:02d}-{slug}.mp3"
    return f"{stem}.c{n:02d}-{slug}.mp3"


def next_audio_version(audios_dir: Path, stem: str) -> int:
    max_v = 0
    if not audios_dir.is_dir():
        return 1
    for p in audios_dir.iterdir():
        if not p.is_file() or p.suffix.lower() != ".mp3":
            continue
        m = VERSION_AUDIO_RE.match(p.name)
        if not m or m.group("stem") != stem:
            continue
        max_v = max(max_v, int(m.group("ver")))
    return max_v + 1


def content_percent_instruction(pct: int) -> str:
    if pct >= 100:
        return ""
    return (
        f"\nIMPORTANTE: sintetiza el contenido para que el guion hablado resultante "
        f"tenga aproximadamente el {pct}% del conteo de palabras del texto original. "
        "Conserva ideas clave y el orden lógico; no inventes hechos."
    )


def deepseek_system_for(pct: int) -> str:
    return PREMIUM_SYSTEM + content_percent_instruction(pct)


def list_stem_audios(audios_dir: Path, stem: str) -> list[Path]:
    out: list[Path] = []
    mono = audios_dir / f"{stem}.mp3"
    if mono.is_file():
        out.append(mono)
    prefix = f"{stem}.c"
    for p in audios_dir.iterdir():
        if p.is_file() and p.name.startswith(prefix) and p.suffix.lower() == ".mp3":
            out.append(p)
    return out


def needs_regen_stem(doc: Path, audios_dir: Path, stem: str, *, premium: bool) -> bool:
    existing = list_stem_audios(audios_dir, stem)
    if not existing:
        return True
    if premium:
        # Premium expects chapter files; a lone legacy stem.mp3 means regenerate.
        chapters = [p for p in existing if f"{stem}.c" in p.name]
        if not chapters:
            return True
        newest_audio = max(p.stat().st_mtime for p in chapters)
    else:
        mono = audios_dir / f"{stem}.mp3"
        if not mono.is_file():
            return True
        newest_audio = mono.stat().st_mtime
    return doc.stat().st_mtime > newest_audio


def clear_stem_audios(audios_dir: Path, stem: str) -> None:
    for p in list_stem_audios(audios_dir, stem):
        try:
            p.unlink()
        except OSError:
            pass


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(". ".join(cells))
    return "\n\n".join(parts)


def extract_pdf_text_layer(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def ocr_pil_image(image, *, name: str, page_label: str = "") -> str:
    """Shared Tesseract path for image files and PDF page renders."""
    import pytesseract
    from PIL import Image, ImageEnhance, ImageOps

    tag = f"{name}{(' ' + page_label) if page_label else ''}"
    log(f"EXTRACT {tag} pct=40 detail=preprocess")
    rgb = image.convert("RGB")
    w, h = rgb.size
    if max(w, h) < 2500:
        rgb = rgb.resize((w * 2, h * 2), Image.Resampling.LANCZOS)
    gray = ImageOps.autocontrast(ImageOps.grayscale(rgb))
    gray = ImageEnhance.Contrast(gray).enhance(1.4)
    log(f"EXTRACT {tag} pct=70 detail=tesseract")
    text = pytesseract.image_to_string(gray, lang="spa+eng").strip()
    log(f"EXTRACT {tag} pct=100 detail=chars={len(text)}")
    return text


def extract_pdf_via_ocr(path: Path) -> str:
    """Render each PDF page and OCR (scanned / image PDFs). Prefer pymupdf; else pdftoppm."""
    name = path.name
    log(f"EXTRACT {name} pct=10 detail=pdf_ocr_start")
    parts: list[str] = []

    try:
        import fitz  # PyMuPDF

        doc = fitz.open(str(path))
        try:
            total = doc.page_count
            for i in range(total):
                page = doc.load_page(i)
                # ~150 DPI matrix (72*2.08 ≈ 150)
                pix = page.get_pixmap(matrix=fitz.Matrix(2.08, 2.08), alpha=False)
                from PIL import Image
                import io

                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text = ocr_pil_image(img, name=name, page_label=f"page={i + 1}/{total}")
                if text:
                    parts.append(text)
        finally:
            doc.close()
        return "\n\n".join(parts)
    except ImportError:
        log(f"EXTRACT {name} detail=pymupdf_missing trying_pdftoppm")
    except Exception as exc:  # noqa: BLE001
        log(f"EXTRACT {name} detail=pymupdf_failed {exc!s}; trying_pdftoppm")

    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError(
            "PDF image OCR needs PyMuPDF (pymupdf) or pdftoppm on PATH"
        )
    with tempfile.TemporaryDirectory(prefix="evoice-pdf-ocr-") as tmp:
        tmp_path = Path(tmp)
        prefix = tmp_path / "page"
        proc = subprocess.run(
            [pdftoppm, "-png", "-r", "150", str(path), str(prefix)],
            capture_output=True,
            check=False,
        )
        if proc.returncode != 0:
            err = proc.stderr.decode("utf-8", errors="replace")[:300]
            raise RuntimeError(f"pdftoppm failed: {err}")
        pages = sorted(tmp_path.glob("page-*.png")) + sorted(tmp_path.glob("page*.png"))
        # pdftoppm names: page-1.png or page1.png depending on version
        if not pages:
            pages = sorted(p for p in tmp_path.iterdir() if p.suffix.lower() == ".png")
        from PIL import Image

        total = len(pages)
        for i, png in enumerate(pages, start=1):
            img = Image.open(png)
            text = ocr_pil_image(img, name=name, page_label=f"page={i}/{total}")
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def extract_pdf(path: Path) -> str:
    """PDF with text layer and/or image pages — always produce text for premium DeepSeek."""
    name = path.name
    log(f"EXTRACT {name} pct=8 detail=pdf_text_layer")
    layer = extract_pdf_text_layer(path)
    if len(layer.strip()) >= PDF_TEXT_MIN_CHARS:
        log(f"EXTRACT {name} pct=100 detail=pdf_text_chars={len(layer)}")
        return layer
    log(
        f"EXTRACT {name} pct=15 detail=pdf_text_sparse chars={len(layer.strip())} "
        "fallback=ocr"
    )
    ocr = extract_pdf_via_ocr(path)
    if layer.strip() and ocr.strip():
        combined = layer.strip() + "\n\n" + ocr.strip()
        log(f"EXTRACT {name} pct=100 detail=pdf_combined_chars={len(combined)}")
        return combined
    if ocr.strip():
        log(f"EXTRACT {name} pct=100 detail=pdf_ocr_chars={len(ocr)}")
        return ocr
    return layer


def extract_image(path: Path) -> str:
    from PIL import Image

    log(f"EXTRACT {path.name} pct=10 detail=open_image")
    image = Image.open(path)
    return ocr_pil_image(image, name=path.name)


def load_doc_text(path: Path) -> str:
    """Extract readable text from every supported modality (spec 044)."""
    ext = path.suffix.lower()
    log(f"EXTRACT {path.name} pct=5 detail=start ext={ext}")
    if ext == ".txt":
        text = extract_txt(path)
        log(f"EXTRACT {path.name} pct=100 detail=chars={len(text)}")
        return text
    if ext == ".docx":
        text = extract_docx(path)
        log(f"EXTRACT {path.name} pct=100 detail=chars={len(text)}")
        return text
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in IMAGE_EXTENSIONS:
        return extract_image(path)
    raise ValueError(f"unsupported extension: {ext}")


def premium_optimize(text: str, name: str, content_percent: int = 100) -> str:
    """DeepSeek chat completions with stream=true + system role; required when DeepSeek runs."""
    key = (os.environ.get("DEEPSEEK_API_KEY") or "").strip()
    if not key:
        raise RuntimeError("DEEPSEEK_API_KEY not configured for premium")
    model = (os.environ.get("DEEPSEEK_MODEL_REASONING") or "deepseek-v4-pro").strip()
    base = (os.environ.get("DEEPSEEK_BASE_URL") or "https://api.deepseek.com").rstrip("/")
    log(f"PREMIUM {name} pct=5 detail=deepseek_stream_start model={model} contentPercent={content_percent}")
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": deepseek_system_for(content_percent)},
            {"role": "user", "content": text[:120000]},
        ],
        "temperature": 0.3,
        "stream": True,
    }
    req = urllib.request.Request(
        f"{base}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {key}",
            "Accept": "text/event-stream",
        },
        method="POST",
    )
    parts: list[str] = []
    last_pct = 5
    try:
        with urllib.request.urlopen(req, timeout=300) as resp:
            while True:
                raw = resp.readline()
                if not raw:
                    break
                line = raw.decode("utf-8", errors="replace").strip()
                if not line:
                    continue
                if not line.startswith("data:"):
                    continue
                data = line[5:].strip()
                if data == "[DONE]":
                    break
                try:
                    chunk = json.loads(data)
                except json.JSONDecodeError:
                    continue
                delta = (
                    chunk.get("choices", [{}])[0]
                    .get("delta", {})
                    .get("content", "")
                )
                if not delta:
                    continue
                parts.append(delta)
                n = sum(len(p) for p in parts)
                # Map growing output length into 10–95 so the UI moves before completion.
                pct = min(95, 10 + (n // 40))
                if pct >= last_pct + 5 or (pct > last_pct and pct >= 90):
                    last_pct = pct
                    snippet = delta.replace("\n", " ")[:80]
                    log(f"PREMIUM {name} pct={pct} detail=chars={n} {snippet}")
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:400]
        raise RuntimeError(f"deepseek HTTP {exc.code}: {detail}") from exc
    content = "".join(parts).strip()
    if not content:
        raise RuntimeError("deepseek returned empty content")
    log(f"PREMIUM {name} pct=100 detail=optimized {len(text)}→{len(content)} chars")
    return content


def vision_image_to_text(image_path: Path, name: str, page_label: str = "") -> str:
    """One page image → DeepSeek Vision (spec 069)."""
    key = (os.environ.get("DEEPSEEK_API_KEY") or "").strip()
    if not key:
        raise RuntimeError("DEEPSEEK_API_KEY not configured for super premium vision")
    model = (
        os.environ.get("DEEPSEEK_MODEL_VISION") or "deepseek-v4-flash-vision-exp"
    ).strip()
    base = (os.environ.get("DEEPSEEK_BASE_URL") or "https://api.deepseek.com").rstrip("/")
    tag = f"{name}{(' ' + page_label) if page_label else ''}"
    log(f"VISION {tag} pct=10 detail=encode")
    import base64
    import gc

    raw = image_path.read_bytes()
    b64 = base64.standard_b64encode(raw).decode("ascii")
    del raw
    mime = "image/jpeg"
    suf = image_path.suffix.lower()
    if suf == ".png":
        mime = "image/png"
    elif suf == ".webp":
        mime = "image/webp"
    elif suf in {".jpg", ".jpeg"}:
        mime = "image/jpeg"
    payload = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": VISION_PAGE_PROMPT},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{b64}"},
                    },
                ],
            }
        ],
        "temperature": 0.1,
        "stream": False,
    }
    # Drop b64 string once embedded so we do not keep two giant copies.
    del b64
    body_bytes = json.dumps(payload).encode("utf-8")
    del payload
    req = urllib.request.Request(
        f"{base}/chat/completions",
        data=body_bytes,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {key}",
        },
        method="POST",
    )
    del body_bytes
    log(f"VISION {tag} pct=40 detail=request model={model}")
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            body = json.loads(resp.read().decode("utf-8", errors="replace"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:400]
        raise RuntimeError(f"vision HTTP {exc.code}: {detail}") from exc
    text = (
        body.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
        or ""
    ).strip()
    del body
    gc.collect()
    log(f"VISION {tag} pct=100 detail=chars={len(text)}")
    return text


def _render_one_pdf_page_pymupdf(
    doc: object,
    page_index: int,
    out_path: Path,
    dpi: int,
) -> None:
    """Render a single 0-based page to JPEG and free the pixmap immediately."""
    import fitz

    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = doc.load_page(page_index).get_pixmap(matrix=mat, alpha=False)
    try:
        try:
            pix.save(str(out_path), jpg_quality=VISION_JPEG_QUALITY)
        except TypeError:
            # Older PyMuPDF: format inferred from .jpg extension
            pix.save(str(out_path))
    finally:
        del pix
        try:
            fitz.TOOLS.store_shrink(100)
        except Exception:  # noqa: BLE001
            pass


def _render_one_pdf_page_pdftoppm(
    pdf_path: Path,
    page_num: int,
    out_path: Path,
    dpi: int,
) -> None:
    """Render one 1-based page via pdftoppm (JPEG)."""
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError(
            "PDF vision needs pymupdf or pdftoppm "
            "(pip install pymupdf OR apt/dnf install poppler-utils)"
        )
    prefix = out_path.parent / f"page-{page_num:04d}"
    proc = subprocess.run(
        [
            pdftoppm,
            "-jpeg",
            "-r",
            str(dpi),
            "-f",
            str(page_num),
            "-l",
            str(page_num),
            "-singlefile",
            str(pdf_path),
            str(prefix),
        ],
        capture_output=True,
        check=False,
    )
    if proc.returncode != 0:
        err = proc.stderr.decode("utf-8", errors="replace")[:300]
        raise RuntimeError(f"pdftoppm failed: {err}")
    produced = prefix.with_suffix(".jpg")
    if not produced.is_file():
        # some builds omit -singlefile naming; glob once
        found = sorted(out_path.parent.glob(f"page-{page_num:04d}*.jpg"))
        if not found:
            raise RuntimeError(f"pdftoppm produced no page {page_num}")
        produced = found[0]
    if produced.resolve() != out_path.resolve():
        produced.replace(out_path)


def extract_via_vision(
    path: Path,
    *,
    checkpoint_path: Path | None = None,
) -> str:
    """Super Premium extract: stream page → Vision → delete (never hold all pages)."""
    import gc

    name = path.name
    ext = path.suffix.lower()
    log(f"VISION {name} pct=5 detail=start ext={ext}")
    parts: list[str] = []
    tmp = Path(tempfile.mkdtemp(prefix="evoice-vision-"))
    try:
        if ext == ".pdf":
            log(f"VISION {name} pct=8 detail=rasterize_stream dpi={VISION_DPI}")
            use_pymupdf = False
            fitz_doc = None
            try:
                import fitz

                fitz_doc = fitz.open(str(path))
                total = int(fitz_doc.page_count)
                use_pymupdf = True
            except ImportError:
                log(f"VISION {name} detail=pymupdf_missing trying_pdftoppm")
                # page count via pdfinfo if available; else probe with pdftoppm
                total = _pdf_page_count_fallback(path)
            except Exception as exc:  # noqa: BLE001
                log(f"VISION {name} detail=pymupdf_failed {exc!s}; trying_pdftoppm")
                if fitz_doc is not None:
                    try:
                        fitz_doc.close()
                    except Exception:  # noqa: BLE001
                        pass
                    fitz_doc = None
                total = _pdf_page_count_fallback(path)
                use_pymupdf = False

            if total <= 0:
                raise RuntimeError("PDF vision: could not determine page count")
            log(
                f"VISION {name} pct=15 detail=stream_pages "
                f"pages={total} engine={'pymupdf' if use_pymupdf else 'pdftoppm'} "
                f"(1 page on disk at a time)"
            )
            try:
                for i in range(1, total + 1):
                    page_file = tmp / f"page-{i:04d}.jpg"
                    try:
                        if use_pymupdf and fitz_doc is not None:
                            _render_one_pdf_page_pymupdf(
                                fitz_doc, i - 1, page_file, VISION_DPI
                            )
                        else:
                            _render_one_pdf_page_pdftoppm(
                                path, i, page_file, VISION_DPI
                            )
                        text = vision_image_to_text(
                            page_file, name, page_label=f"page={i}/{total}"
                        )
                        if text:
                            parts.append(text)
                    finally:
                        try:
                            page_file.unlink(missing_ok=True)
                        except TypeError:
                            # py<3.8 compat — not expected on EC2
                            if page_file.exists():
                                page_file.unlink()
                        # clear any leftover pdftoppm siblings
                        for leftover in tmp.glob(f"page-{i:04d}*"):
                            try:
                                leftover.unlink()
                            except OSError:
                                pass
                    if checkpoint_path is not None and (
                        i % VISION_CHECKPOINT_EVERY == 0 or i == total
                    ):
                        checkpoint_path.write_text(
                            "\n\n".join(parts), encoding="utf-8"
                        )
                        log(
                            f"VISION {name} detail=checkpoint "
                            f"pages={i}/{total} chars={sum(len(p) for p in parts)}"
                        )
                    if i % 5 == 0:
                        gc.collect()
            finally:
                if fitz_doc is not None:
                    fitz_doc.close()
        elif ext in IMAGE_EXTENSIONS:
            text = vision_image_to_text(path, name)
            if text:
                parts.append(text)
            if checkpoint_path is not None and parts:
                checkpoint_path.write_text("\n\n".join(parts), encoding="utf-8")
        else:
            raise ValueError(f"vision unsupported: {ext}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        gc.collect()
    joined = "\n\n".join(parts)
    if checkpoint_path is not None:
        checkpoint_path.write_text(joined, encoding="utf-8")
    log(f"VISION {name} pct=100 detail=chars={len(joined)}")
    return joined


def _pdf_page_count_fallback(path: Path) -> int:
    """Best-effort page count when PyMuPDF is unavailable."""
    pdfinfo = shutil.which("pdfinfo")
    if pdfinfo:
        proc = subprocess.run(
            [pdfinfo, str(path)], capture_output=True, check=False, text=True
        )
        if proc.returncode == 0:
            for line in (proc.stdout or "").splitlines():
                if line.lower().startswith("pages:"):
                    try:
                        return int(line.split(":", 1)[1].strip())
                    except ValueError:
                        break
    # Last resort: pypdf
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path)).pages)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"PDF page count failed: {exc}") from exc


def super_allowed(path: Path) -> bool:
    ext = path.suffix.lower()
    return ext in {".pdf", ".docx"} | IMAGE_EXTENSIONS


def find_ffmpeg() -> str:
    which = shutil.which("ffmpeg")
    if not which:
        raise FileNotFoundError("ffmpeg not found on PATH")
    return which


def text_to_wav_piper(text: str, wav_path: Path) -> None:
    piper = shutil.which("piper")
    if not piper:
        raise FileNotFoundError("piper not found")
    model = Path(__file__).resolve().parent / "models" / "es_ES-sharvard-medium.onnx"
    env_model = Path(os.environ.get("EVOICE_PIPER_MODEL", "")).expanduser()
    if env_model.is_file():
        model = env_model
    if not model.is_file():
        raise FileNotFoundError(f"piper model missing: {model}")
    proc = subprocess.run(
        [piper, "--model", str(model), "--output_file", str(wav_path)],
        input=text.encode("utf-8"),
        capture_output=True,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.decode("utf-8", errors="replace") or "piper failed")


def text_to_wav_espeak(text: str, wav_path: Path) -> None:
    espeak = shutil.which("espeak-ng") or shutil.which("espeak")
    if not espeak:
        raise FileNotFoundError("espeak-ng not found")
    proc = subprocess.run(
        [espeak, "-v", "es", "-w", str(wav_path), text],
        capture_output=True,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.decode("utf-8", errors="replace") or "espeak failed")


def wav_to_mp3(wav_path: Path, mp3_path: Path) -> None:
    ffmpeg = find_ffmpeg()
    log(f"FFMPEG {mp3_path.name} pct=50 detail=encoding")
    proc = subprocess.run(
        [
            ffmpeg,
            "-y",
            "-i",
            str(wav_path),
            "-codec:a",
            "libmp3lame",
            "-qscale:a",
            "4",
            str(mp3_path),
        ],
        capture_output=True,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.decode("utf-8", errors="replace")[-400:] or "ffmpeg failed")
    log(f"FFMPEG {mp3_path.name} pct=100 detail=ok")


def chunk_text(text: str, max_chars: int = 800) -> list[str]:
    text = text.strip()
    if not text:
        return []
    parts: list[str] = []
    buf: list[str] = []
    size = 0
    for para in text.replace("\r", "").split("\n"):
        para = para.strip()
        if not para:
            continue
        if size + len(para) + 1 > max_chars and buf:
            parts.append(" ".join(buf))
            buf = [para]
            size = len(para)
        else:
            buf.append(para)
            size += len(para) + 1
    if buf:
        parts.append(" ".join(buf))
    return parts or [text]


def text_to_mp3(text: str, mp3_path: Path, name: str, tmp_parent: Path | None = None) -> None:
    text = text.strip()
    if not text:
        raise ValueError("empty text")
    parent = tmp_parent if tmp_parent is not None else mp3_path.parent
    parent.mkdir(parents=True, exist_ok=True)
    chunks = chunk_text(text)
    with tempfile.TemporaryDirectory(prefix="evoice-tts-", dir=str(parent)) as tmp:
        tmp_path = Path(tmp)
        wav_parts: list[Path] = []
        for i, chunk in enumerate(chunks, start=1):
            pct = int(100 * (i - 1) / max(len(chunks), 1))
            log(f"TTS {name} pct={pct} detail=chunk {i}/{len(chunks)}")
            wav = tmp_path / f"part-{i:04d}.wav"
            try:
                text_to_wav_piper(chunk, wav)
                if i == 1:
                    log(f"TTS {name} detail=engine=piper")
            except Exception as piper_err:  # noqa: BLE001
                if i == 1:
                    log(f"TTS {name} detail=piper_fail ({piper_err}); espeak-ng")
                text_to_wav_espeak(chunk, wav)
            wav_parts.append(wav)
            log(f"TTS {name} pct={int(100 * i / max(len(chunks), 1))} detail=chunk_done")
        if len(wav_parts) == 1:
            wav_to_mp3(wav_parts[0], mp3_path)
            return
        list_file = tmp_path / "concat.txt"
        list_file.write_text(
            "\n".join(f"file '{p.resolve().as_posix()}'" for p in wav_parts),
            encoding="utf-8",
        )
        merged = tmp_path / "merged.wav"
        ffmpeg = find_ffmpeg()
        proc = subprocess.run(
            [ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", str(merged)],
            capture_output=True,
            check=False,
        )
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.decode("utf-8", errors="replace")[-400:] or "ffmpeg concat failed")
        wav_to_mp3(merged, mp3_path)


def sync_project(
    project_dir: Path,
    only: set[str] | None = None,
    *,
    mode: str = "standard",
    content_percent: int = 100,
) -> dict[str, int]:
    mode = (mode or "standard").strip().lower()
    if mode in {"super", "superpremium", "super-premium"}:
        mode = "super_premium"
    if mode not in {"standard", "premium", "super_premium"}:
        mode = "standard"
    if content_percent not in {100, 75, 50, 25, 10, 5}:
        content_percent = 100
    use_deepseek = mode in {"premium", "super_premium"} or content_percent < 100
    docs_dir = project_dir / "docs"
    audios_dir = project_dir / "audios"
    audios_dir.mkdir(parents=True, exist_ok=True)
    docs = sorted(
        p
        for p in docs_dir.iterdir()
        if p.is_file()
        and p.name != ".keep"
        and not p.name.endswith(".premium.txt")
        and not p.name.endswith(".vision.txt")
        and p.suffix.lower() in DOC_EXTENSIONS
        and (not only or p.name in only)
    )
    stats = {"docs": len(docs), "generated": 0, "skipped": 0, "failed": 0}
    log(f"STEP convert docs={len(docs)} mode={mode} contentPercent={content_percent}")
    if not docs:
        log("No convertible files in docs/")
        return stats
    for idx, doc in enumerate(docs, start=1):
        stem = doc.stem
        log(f"FILE {doc.name} state=active")
        log(f"STEP convert doc={idx}/{len(docs)} file={doc.name}")
        ver = next_audio_version(audios_dir, stem)
        log(f"gen   {doc.name} (version=v{ver})")
        try:
            if mode == "super_premium":
                if doc.suffix.lower() == ".txt":
                    raise ValueError("super_premium not allowed for .txt")
                if not super_allowed(doc):
                    raise ValueError(f"super_premium unsupported for {doc.suffix}")
                if doc.suffix.lower() == ".docx":
                    text = load_doc_text(doc)
                else:
                    vision_path = docs_dir / f"{stem}.v{ver}.vision.txt"
                    text = extract_via_vision(doc, checkpoint_path=vision_path)
                    log(f"VISION {doc.name} detail=wrote {vision_path.name}")
            else:
                text = load_doc_text(doc)
            if not text.strip():
                raise ValueError("No readable text extracted")
            if use_deepseek:
                log(
                    f"PREMIUM {doc.name} detail=system_role_prep modality={doc.suffix.lower()} "
                    f"mode={mode}"
                )
                text = premium_optimize(text, doc.name, content_percent)
                premium_path = docs_dir / f"{stem}.v{ver}.premium.txt"
                premium_path.write_text(text, encoding="utf-8")
                log(f"PREMIUM {doc.name} detail=wrote {premium_path.name}")
                chapters = parse_chapters(text)
                if not chapters:
                    raise ValueError("deepseek produced no chapters")
                log(f"PREMIUM {doc.name} detail=chapters={len(chapters)}")
                for i, (n, title, body) in enumerate(chapters, start=1):
                    mp3_name = chapter_mp3_name(stem, n, title, version=ver)
                    mp3 = audios_dir / mp3_name
                    pct = int(100 * (i - 1) / max(len(chapters), 1))
                    log(f"TTS {doc.name} pct={pct} detail=chapter {n}/{len(chapters)} {title}")
                    text_to_mp3(body, mp3, doc.name, tmp_parent=project_dir)
                    log(f"ok     {doc.name} -> {mp3_name}")
            else:
                mp3_name = f"{stem}.v{ver}.mp3"
                mp3 = audios_dir / mp3_name
                text_to_mp3(text, mp3, doc.name, tmp_parent=project_dir)
                log(f"ok     {doc.name} -> {mp3_name}")
            stats["generated"] += 1
            log(f"FILE {doc.name} state=done")
        except Exception as exc:  # noqa: BLE001
            stats["failed"] += 1
            log(f"FAIL  {doc.name}: {exc}")
            log(f"FILE {doc.name} state=failed")
    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-dir", type=Path, required=True)
    parser.add_argument(
        "--only",
        action="append",
        default=[],
        help="Only convert this basename (repeatable). Default: all convertible docs.",
    )
    parser.add_argument(
        "--premium",
        action="store_true",
        help="Legacy: same as --mode premium.",
    )
    parser.add_argument(
        "--mode",
        default="",
        help="standard | premium | super_premium (spec 069).",
    )
    parser.add_argument(
        "--content-percent",
        type=int,
        default=100,
        help="Target spoken length as %% of original word count (100/75/50/25/10/5).",
    )
    args = parser.parse_args()
    project_dir = args.project_dir.resolve()
    if not (project_dir / "docs").is_dir():
        log(f"docs/ missing under {project_dir}")
        return 1
    only = {n for n in args.only if n} or None
    mode = (args.mode or "").strip().lower()
    if not mode:
        mode = "premium" if args.premium else "standard"
    stats = sync_project(
        project_dir,
        only=only,
        mode=mode,
        content_percent=int(args.content_percent),
    )
    log(
        f"STATS docs={stats['docs']} generated={stats['generated']} "
        f"skipped={stats['skipped']} failed={stats['failed']}"
    )
    return 0 if stats["failed"] == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
